import csv
import io
import json
import openpyxl

class SimpleDataFrame:
    """
    Lightweight, pure-Python DataFrame replacement for Cloudflare Workers.
    Provides dict-like and record-based access without the heavy 100MB+ pandas/numpy bundle.
    """
    def __init__(self, columns=None, records=None):
        self.columns = list(columns) if columns else []
        self._records = list(records) if records else []

    def fillna(self, val=''):
        cleaned = []
        for row in self._records:
            cleaned.append({k: (val if v is None or v == '' else v) for k, v in row.items()})
        return SimpleDataFrame(self.columns, cleaned)

    def iterrows(self):
        for idx, row in enumerate(self._records):
            yield idx, row

    def head(self, n=5):
        return SimpleDataFrame(self.columns, self._records[:n])

    def to_dict(self, orient='records'):
        if orient == 'records':
            return [dict(r) for r in self._records]
        return [dict(r) for r in self._records]

    def get(self, key, default=None):
        return self._records if key == 'records' else default

    def __len__(self):
        return len(self._records)

    def __iter__(self):
        return iter(self._records)


DataFrame = SimpleDataFrame


def read_excel(file_obj, sheet_name=None):
    """
    Reads an Excel file using openpyxl and returns SimpleDataFrame objects.
    """
    if hasattr(file_obj, 'seek'):
        file_obj.seek(0)
    wb = openpyxl.load_workbook(file_obj, data_only=True)
    sheets = {}

    for title in wb.sheetnames:
        sheet = wb[title]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            sheets[title] = SimpleDataFrame([], [])
            continue

        # Extract headers from row 1
        headers = [str(cell).strip() if cell is not None else f"col_{i}" for i, cell in enumerate(rows[0])]
        records = []

        for row in rows[1:]:
            if all(c is None for c in row):
                continue
            record = {}
            for i, h in enumerate(headers):
                val = row[i] if i < len(row) else ''
                record[h] = '' if val is None else val
            records.append(record)

        sheets[title] = SimpleDataFrame(headers, records)

    if sheet_name is None:
        return sheets
    return sheets.get(sheet_name, SimpleDataFrame([], []))


def read_csv(file_obj):
    """
    Reads CSV content and returns a SimpleDataFrame.
    """
    if hasattr(file_obj, 'seek'):
        file_obj.seek(0)
    raw = file_obj.read()
    if isinstance(raw, bytes):
        text = raw.decode('utf-8', errors='ignore')
    else:
        text = str(raw)

    reader = csv.DictReader(io.StringIO(text))
    records = [dict(r) for r in reader]
    headers = list(reader.fieldnames) if reader.fieldnames else []
    return SimpleDataFrame(headers, records)


def read_json(file_obj):
    """
    Reads JSON content and returns a SimpleDataFrame.
    """
    if hasattr(file_obj, 'seek'):
        file_obj.seek(0)
    raw = file_obj.read()
    if isinstance(raw, bytes):
        text = raw.decode('utf-8', errors='ignore')
    else:
        text = str(raw)

    data = json.loads(text)
    if isinstance(data, list) and len(data) > 0 and isinstance(data[0], dict):
        headers = list(data[0].keys())
        return SimpleDataFrame(headers, data)
    return SimpleDataFrame([], [])
